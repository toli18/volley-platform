#!/usr/bin/env python3
"""Build the official Appendix 12 DOCX blank (structure from the municipality form extract).

Run: python scripts/build_troyan_hall_docx_template.py
Writes: backend/app/static/troyan-hall/prilozhenie-12.docx
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "backend" / "app" / "static" / "troyan-hall" / "prilozhenie-12.docx"
LOGO = ROOT / "backend" / "app" / "static" / "troyan-hall" / "obshtina-troyan-logo.png"


def _set_run_font(run, *, size=11, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _p(doc, text="", *, size=11, bold=False, italic=False, align=None, space_after=2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        _set_run_font(r, size=size, bold=bold, italic=italic)
    return p


def _add_runs(p, parts):
    """parts: list of (text, size, bold, italic)"""
    for text, size, bold, italic in parts:
        r = p.add_run(text)
        _set_run_font(r, size=size, bold=bold, italic=italic)


def build() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # --- Header: logo + ОБЩИНА ТРОЯН | Приложение / Вх. / Дата ---
    header = doc.add_table(rows=1, cols=2)
    header.autofit = True
    left, right = header.rows[0].cells

    if LOGO.is_file():
        lp = left.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        run = lp.add_run()
        run.add_picture(str(LOGO), width=Cm(1.4))
        p2 = left.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(0)
        _add_runs(p2, [("ОБЩИНА ТРОЯН", 12, True, False)])
        p3 = left.add_paragraph()
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(0)
        _add_runs(p3, [("№ .............................................", 10, False, False)])
    else:
        _add_runs(left.paragraphs[0], [("ОБЩИНА ТРОЯН", 12, True, False)])

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_after = Pt(0)
    _add_runs(rp, [("Приложение № 12", 11, True, False)])
    r2 = right.add_paragraph()
    r2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2.paragraph_format.space_before = Pt(0)
    r2.paragraph_format.space_after = Pt(0)
    _add_runs(r2, [("(ново, Решение № 965/17.09.2026 г.)", 9, False, True)])
    r3 = right.add_paragraph()
    r3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r3.paragraph_format.space_before = Pt(4)
    r3.paragraph_format.space_after = Pt(0)
    _add_runs(r3, [("Вх. № .............................................", 10, False, False)])
    r4 = right.add_paragraph()
    r4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r4.paragraph_format.space_before = Pt(0)
    r4.paragraph_format.space_after = Pt(0)
    _add_runs(r4, [("Дата ...............................................", 10, False, False)])

    _p(doc, "", space_after=6)

    _p(doc, "ЗАЯВЛЕНИЕ", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(
        doc,
        "ползване на спортни зали и помещения в",
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=0,
    )
    _p(
        doc,
        "Многофункционална спортна зала „Троян“",
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )

    _p(doc, "от {{ORG_NAME}}", size=11, space_after=0)
    _p(
        doc,
        "(посочете трите имена за физическото лице или наименованието на юридическото лице)",
        size=9,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )

    p = _p(doc, "", size=11, space_after=0)
    _add_runs(
        p,
        [
            (
                "ЕГН/ЕИК {{EIK}}, постоянен/настоящ адрес или адрес на управление на юридическото лице: "
                "гр./с. {{CITY}}, община {{MUNICIPALITY}}, област {{OBLAST}}, ул. (ж.к.) {{STREET}}, "
                "тел.: {{PHONE}} представляван/а от {{REP_NAME}}",
                11,
                False,
                False,
            )
        ],
    )
    _p(
        doc,
        "(трите имена на упълномощен или законен представител на юридическото лице)",
        size=9,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    _p(doc, "e-mail: {{EMAIL}}", size=11, space_after=8)

    _p(doc, "УВАЖАЕМА ГОСПОЖО КМЕТ,", size=11, bold=True, space_after=6)
    _p(
        doc,
        "Заявявам желание за предоставяне за ползване на следното помещение в Многофункционална спортна зала „Троян“",
        size=11,
        space_after=4,
    )
    _p(
        doc,
        "{{CHK_HALL_SPORT}} Спортна зала;   {{CHK_HALL_TRAIN}} Тренировъчна зала;  {{CHK_HALL_CONF}} Конферентна зала (30 места)",
        size=11,
        space_after=6,
    )
    _p(doc, "за провеждане на   {{EVENT_TITLE}}", size=11, space_after=0)
    _p(
        doc,
        "(наименование на събитието)",
        size=9,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    _p(doc, "вид на събитието", size=11, space_after=2)
    _p(doc, "{{CHK_KIND_TRAIN}} Тренировки;", size=11, space_after=0)
    _p(doc, "{{CHK_KIND_TOURN}} Спортни турнири/състезания;", size=11, space_after=0)
    _p(doc, "{{CHK_KIND_OTHER}} Образователни, културни или социални събития и дейности;", size=11, space_after=6)
    _p(doc, "дати и часове за ползване:", size=11, space_after=4)

    # Schedule table — 10 data rows as on the blank
    sched = doc.add_table(rows=11, cols=4)
    sched.style = "Table Grid"
    headers = [
        "Дата",
        "Начален час",
        "Краен час",
        "Целодневно ползване\n(до 8 часа в диапазона\n08:00 ч. – 22:00 ч.)",
    ]
    for i, h in enumerate(headers):
        cell = sched.rows[0].cells[i]
        cell.text = ""
        cp = cell.paragraphs[0]
        r = cp.add_run(h)
        _set_run_font(r, size=9, bold=True)
    for ri in range(1, 11):
        for ci in range(4):
            cell = sched.rows[ri].cells[ci]
            cell.text = ""
            cp = cell.paragraphs[0]
            txt = "☐" if ci == 3 else ""
            r = cp.add_run(txt)
            _set_run_font(r, size=10)

    _p(doc, "", space_after=6)
    _p(doc, "За провеждането на събитието е необходимо:", size=11, space_after=4)
    _p(
        doc,
        "{{CHK_EQ_BASKET}} Оборудване и спортни съоръжения - Баскетбол;   "
        "{{CHK_EQ_VOLLEY}} Оборудване и спортни съоръжения – Волейбол;   "
        "{{CHK_EQ_TENNIS}} Оборудване и спортни съоръжения – Тенис;    "
        "{{CHK_EQ_VINYL}} Монтиране на предпазна винилова настилка;   "
        "{{CHK_EQ_SOUND}} Озвучителна система (оповестителна);   "
        "{{CHK_EQ_SCREEN}} Лед екран, табла за резултати;",
        size=10,
        space_after=6,
    )
    _p(
        doc,
        "Изисква се преустройство на залата:  {{CHK_REBUILD_YES}} ДА;  {{CHK_REBUILD_NO}} НЕ",
        size=11,
        space_after=0,
    )
    _p(
        doc,
        "(поставяне на предпазна винилова настилка, внасяне, монтаж и демонтаж на сцена, подиуми и сценично оборудване)",
        size=9,
        italic=True,
        space_after=6,
    )
    _p(doc, "Други бележки: ........................................................................", size=11, space_after=0)
    _p(doc, "........................................................................................", size=11, space_after=0)
    _p(doc, "........................................................................................", size=11, space_after=8)

    _p(
        doc,
        "Запознат/а съм, че съгласно Наредба № 7 за определянето и администрирането на местните такси и цени на услуги на територията на община Троян, цената на услуга с ДДС е:",
        size=10,
        space_after=4,
    )
    _p(doc, "За спортна зала:", size=11, bold=True, space_after=2)

    def price_table(rows):
        t = doc.add_table(rows=len(rows), cols=2)
        t.style = "Table Grid"
        for ri, (a, b) in enumerate(rows):
            for ci, val in enumerate((a, b)):
                cell = t.rows[ri].cells[ci]
                cell.text = ""
                cp = cell.paragraphs[0]
                r = cp.add_run(val)
                _set_run_font(r, size=9, bold=(ri == 0))
        _p(doc, "", space_after=4)

    price_table(
        [
            ("Наименование на услугата", "цена в евро"),
            (
                "1.1. Почасово ползване за тренировъчна дейност и провеждане на състезания и турнири от спортни клубове, регистрирани и осъществяващи дейност на територията на община Троян",
                "8,00 евро/час",
            ),
            (
                "1.2. Почасово ползване за тренировъчна дейност на външни спортни клубове и организации",
                "25,00 евро/час",
            ),
            ("1.3. Дневно ползване (до 8 часа в диапазона 08:00 ч. – 22:00 ч.)", ""),
            (
                "1.3.1. за турнири и спортни прояви с публика, организирани от външни клубове – 240,00 евро",
                "240 евро",
            ),
            ("1.3.2.  за събития с ниско натоварване", "240 евро"),
            ("1.3.3.  за събития със средно натоварване", "360 евро"),
            ("1.3.4.  за събития с високо натоварване", "480 евро"),
        ]
    )
    _p(doc, "За тренировъчна зала:", size=11, bold=True, space_after=2)
    price_table(
        [
            ("Наименование на услугата", "цена в евро"),
            (
                "2.1. Почасово ползване за тренировъчна дейност и провеждане на състезания и турнири от спортни клубове, регистрирани и осъществяващи дейност на територията на община Троян",
                "5,00 евро/час",
            ),
            (
                "2.2. Почасово ползване за тренировъчна дейност и провеждане на състезания и турнири на външни спортни клубове и организации",
                "15,00 евро/час",
            ),
        ]
    )
    _p(doc, "За конферентна зала:", size=11, bold=True, space_after=2)
    price_table(
        [
            ("Наименование на услугата", "цена в евро"),
            (
                "3.1. Почасово ползване на конферентна зала в Многофункционална спортна зала „Троян“",
                "15,00 евро/час",
            ),
        ]
    )

    _p(
        doc,
        "Дължимата сума се заплаща до 3 работни дни след получаване на Разрешение за ползване на:",
        size=10,
        space_after=2,
    )
    _p(
        doc,
        "1. Касата на Община Троян – административна сграда на адрес Троян, пл. „Възраждане“ № 1",
        size=10,
        space_after=0,
    )
    _p(doc, "2. С ПОС терминал в отдел „Услуги на гражданите“", size=10, space_after=0)
    _p(doc, "3. По банков път на сметка:", size=10, space_after=0)
    _p(doc, "IBAN: BG92UBBS81558420008108", size=10, space_after=0)
    _p(doc, "BIC: UBBSBGSF", size=10, space_after=0)
    _p(doc, "Банка: „Обединена Българска Банка“ АД", size=10, space_after=6)
    _p(
        doc,
        "* Забележка: При невнасяне на дължимите такси в регламентирания срок, заявлението се счита оттеглено.",
        size=9,
        italic=True,
        space_after=10,
    )

    # Signature block with tab
    sig = doc.add_paragraph()
    sig.paragraph_format.space_after = Pt(0)
    tab_stops = sig.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Twips(9000), WD_TAB_ALIGNMENT.RIGHT)
    _add_runs(
        sig,
        [
            ("Дата: {{TODAY}} г.", 11, False, False),
            ("\t", 11, False, False),
            ("С уважение: ..............................................", 11, False, False),
        ],
    )
    sig2 = doc.add_paragraph()
    sig2.paragraph_format.space_before = Pt(0)
    tab_stops2 = sig2.paragraph_format.tab_stops
    tab_stops2.add_tab_stop(Twips(9000), WD_TAB_ALIGNMENT.RIGHT)
    _add_runs(
        sig2,
        [
            ("Град Троян", 11, False, False),
            ("\t", 11, False, False),
            ("(подпис)", 10, False, False),
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path} ({path.stat().st_size} bytes)")
