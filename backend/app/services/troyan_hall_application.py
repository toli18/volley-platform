"""Месечно заявление за МФЗ „Троян“ (Приложение 12) — само ВК Троян Волей."""

from __future__ import annotations

import calendar
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from app.models import Club, TrainingScheduleException, TrainingScheduleRule, User
from app.services.school_excuse_note import is_troyan_volley_club


MONTH_BG = (
    "",
    "януари",
    "февруари",
    "март",
    "април",
    "май",
    "юни",
    "юли",
    "август",
    "септември",
    "октомври",
    "ноември",
    "декември",
)

TROYAN_ORG_NAME = "Волейболен клуб Троян Волей"
TROYAN_EIK = "205338036"
TROYAN_CITY = "Троян"
TROYAN_MUNICIPALITY = "Троян"
TROYAN_OBLAST = "Ловеч"
TROYAN_STREET = "Захарий Зограф № 31"
LOCATION_NEEDLE = "троян"


def event_title_for_month(year: int, month: int) -> str:
    return f"Тренировки по волейбол – месец {MONTH_BG[int(month)]} {int(year)}"


def event_title_for_period(date_from: date, date_to: date) -> str:
    if date_from.year == date_to.year and date_from.month == date_to.month:
        last = calendar.monthrange(date_from.year, date_from.month)[1]
        if date_from.day == 1 and date_to.day == last:
            return event_title_for_month(date_from.year, date_from.month)
        if date_from.day == date_to.day:
            return (
                f"Тренировки по волейбол – {date_from.day:02d}.{date_from.month:02d}.{date_from.year}"
            )
        return (
            f"Тренировки по волейбол – {date_from.day:02d}–{date_to.day:02d} "
            f"{MONTH_BG[date_from.month]} {date_from.year}"
        )
    return (
        f"Тренировки по волейбол – {date_from.day:02d}.{date_from.month:02d}.{date_from.year}"
        f" – {date_to.day:02d}.{date_to.month:02d}.{date_to.year}"
    )


def _parse_iso_date(value: str | None, label: str) -> date | None:
    raw = (value or "").strip()
    if not raw:
        return None
    try:
        y, m, d = raw.split("-")
        return date(int(y), int(m), int(d))
    except Exception as exc:
        raise ValueError(f"Невалидна дата за {label}") from exc


def resolve_period(
    year: int | None,
    month: int | None,
    date_from: str | None = None,
    date_to: str | None = None,
) -> tuple[date, date]:
    """Период от–до. Ако липсват дати — целият месец (year/month или текущ)."""
    d_from = _parse_iso_date(date_from, "date_from")
    d_to = _parse_iso_date(date_to, "date_to")
    today = date.today()
    y = int(year or (d_from.year if d_from else today.year))
    m = int(month or (d_from.month if d_from else today.month))
    if m < 1 or m > 12:
        raise ValueError("month must be 1–12")
    last = calendar.monthrange(y, m)[1]
    month_start = date(y, m, 1)
    month_end = date(y, m, last)
    if d_from is None:
        d_from = month_start
    if d_to is None:
        d_to = month_end
    if d_to < d_from:
        raise ValueError("Крайната дата трябва да е след началната")
    if (d_to - d_from).days > 62:
        raise ValueError("Периодът е твърде дълъг (макс. 62 дни)")
    return d_from, d_to


def legal_defaults(club: Club | None, current_user: User | None = None) -> dict[str, Any]:
    chairman = ""
    if club:
        chairman = (
            (getattr(club, "school_excuse_chairman_name", None) or "").strip()
            or (club.contact_name or "").strip()
        )
    if not chairman and current_user:
        chairman = (current_user.name or "").strip()
    phone = (club.contact_phone or "").strip() if club else ""
    email = (club.contact_email or "").strip() if club else ""
    eik = (club.bulstat or "").strip() if club else ""
    if not eik:
        eik = TROYAN_EIK
    org = ""
    if club:
        org = (club.full_name or club.name or "").strip()
    if not org or "троян" not in org.lower():
        org = TROYAN_ORG_NAME
    return {
        "org_name": org,
        "eik": eik,
        "city": TROYAN_CITY,
        "municipality": TROYAN_MUNICIPALITY,
        "oblast": TROYAN_OBLAST,
        "street": TROYAN_STREET,
        "representative_name": chairman,
        "phone": phone,
        "email": email,
        "hall_type": "sport",  # спортна зала
        "event_kind": "training",
        "equipment_volleyball": True,
        "rebuild_hall": False,
    }


def _norm_location(value: str | None) -> str:
    return " ".join(str(value or "").split()).strip()


def _is_troyan_location(location: str | None) -> bool:
    return LOCATION_NEEDLE in _norm_location(location).lower()


def _time_to_minutes(hhmm: str) -> int:
    parts = str(hhmm or "0:0").split(":")
    h = int(parts[0] or 0)
    m = int(parts[1] or 0) if len(parts) > 1 else 0
    return h * 60 + m


def _minutes_to_time(total: int) -> str:
    h, m = divmod(int(total), 60)
    return f"{h:02d}:{m:02d}"


def build_month_slots(
    db: Session,
    club_id: int,
    year: int | None = None,
    month: int | None = None,
    *,
    date_from: str | None = None,
    date_to: str | None = None,
) -> list[dict[str, Any]]:
    """Един ред на ден: обединени часове за всички тренировки с локация „Троян“."""
    d0, d1 = resolve_period(year, month, date_from, date_to)

    rules = (
        db.query(TrainingScheduleRule)
        .filter(
            TrainingScheduleRule.club_id == int(club_id),
            TrainingScheduleRule.is_active.is_(True),
            TrainingScheduleRule.effective_from <= d1.isoformat(),
        )
        .filter(
            (TrainingScheduleRule.effective_to.is_(None))
            | (TrainingScheduleRule.effective_to >= d0.isoformat())
        )
        .all()
    )
    if not rules:
        return []

    rule_ids = [r.id for r in rules]
    exc_rows = (
        db.query(TrainingScheduleException)
        .filter(
            TrainingScheduleException.rule_id.in_(rule_ids),
            TrainingScheduleException.date >= d0.isoformat(),
            TrainingScheduleException.date <= d1.isoformat(),
        )
        .all()
    )
    exc_by_key = {(e.rule_id, e.date): e for e in exc_rows}

    by_date: dict[str, dict[str, Any]] = {}
    for day_idx in range((d1 - d0).days + 1):
        cur = d0 + timedelta(days=day_idx)
        weekday = cur.weekday()
        cur_s = cur.isoformat()
        for r in rules:
            if int(r.weekday) != int(weekday):
                continue
            if r.effective_from > cur_s:
                continue
            if r.effective_to and r.effective_to < cur_s:
                continue
            exc = exc_by_key.get((r.id, cur_s))
            if exc and exc.kind == "cancelled":
                continue
            location_v = (
                _norm_location(exc.location)
                if exc and exc.kind == "override" and exc.location
                else _norm_location(r.location)
            )
            if not _is_troyan_location(location_v):
                continue
            start_v = exc.start_time if exc and exc.kind == "override" and exc.start_time else r.start_time
            end_v = exc.end_time if exc and exc.kind == "override" and exc.end_time else r.end_time
            slot = by_date.get(cur_s)
            start_m = _time_to_minutes(start_v)
            end_m = _time_to_minutes(end_v)
            if slot is None:
                by_date[cur_s] = {
                    "date": cur_s,
                    "start_time": start_v,
                    "end_time": end_v,
                    "_start_m": start_m,
                    "_end_m": end_m,
                    "session_count": 1,
                    "locations": [location_v] if location_v else [],
                }
            else:
                slot["_start_m"] = min(slot["_start_m"], start_m)
                slot["_end_m"] = max(slot["_end_m"], end_m)
                slot["start_time"] = _minutes_to_time(slot["_start_m"])
                slot["end_time"] = _minutes_to_time(slot["_end_m"])
                slot["session_count"] = int(slot["session_count"]) + 1
                if location_v and location_v not in slot["locations"]:
                    slot["locations"].append(location_v)

    out: list[dict[str, Any]] = []
    for key in sorted(by_date.keys()):
        row = by_date[key]
        hours = max(0.0, (row["_end_m"] - row["_start_m"]) / 60.0)
        out.append(
            {
                "date": row["date"],
                "date_bg": f"{row['date'][8:10]}.{row['date'][5:7]}.{row['date'][0:4]}",
                "start_time": row["start_time"],
                "end_time": row["end_time"],
                "hours": round(hours, 2),
                "session_count": row["session_count"],
                "locations": row["locations"],
                "all_day": False,
            }
        )
    return out


def preview_payload(
    db: Session,
    club: Club,
    year: int | None = None,
    month: int | None = None,
    current_user: User | None = None,
    *,
    date_from: str | None = None,
    date_to: str | None = None,
) -> dict[str, Any]:
    if not is_troyan_volley_club(club):
        raise PermissionError("Заявлението за зала Троян е само за ВК Троян Волей")
    d0, d1 = resolve_period(year, month, date_from, date_to)
    slots = build_month_slots(
        db,
        club.id,
        d0.year,
        d0.month,
        date_from=d0.isoformat(),
        date_to=d1.isoformat(),
    )
    defaults = legal_defaults(club, current_user)
    total_hours = round(sum(float(s["hours"]) for s in slots), 2)
    return {
        "year": d0.year,
        "month": d0.month,
        "month_label": MONTH_BG[d0.month],
        "date_from": d0.isoformat(),
        "date_to": d1.isoformat(),
        "event_title": event_title_for_period(d0, d1),
        "count": len(slots),
        "total_hours": total_hours,
        "slots": slots,
        "defaults": defaults,
        "hall_label": "Спортна зала",
        "rule": "Само тренировки с място, съдържащо „Троян“. По 1 ред на ден с обединени часове.",
    }



def _template_docx_path() -> Path:
    path = Path(__file__).resolve().parent.parent / "static" / "troyan-hall" / "prilozhenie-12.docx"
    if not path.is_file():
        raise FileNotFoundError("Липсва бланката Приложение 12 (prilozhenie-12.docx)")
    return path


def _chk(checked: bool) -> str:
    return "☑" if checked else "☐"


def _replace_in_paragraph(paragraph, mapping: dict[str, str]) -> None:
    """Replace placeholders per-run so Times New Roman / bold / italic stay intact."""
    if not paragraph.runs:
        return
    for run in paragraph.runs:
        text = run.text or ""
        if not text:
            continue
        new = text
        changed = False
        for key, val in mapping.items():
            if key in new:
                new = new.replace(key, val)
                changed = True
        if changed:
            run.text = new


def _iter_cell_paragraphs(cell):
    for p in cell.paragraphs:
        yield p
    for nested in cell.tables:
        for row in nested.rows:
            for nested_cell in row.cells:
                yield from _iter_cell_paragraphs(nested_cell)


def _iter_all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_cell_paragraphs(cell)


def build_troyan_hall_application_docx(
    *,
    year: int,
    month: int,
    slots: list[dict[str, Any]],
    fields: dict[str, Any],
) -> bytes:
    """Fill the official Appendix 12 DOCX blank — layout unchanged, only field values."""
    from copy import deepcopy

    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    template = _template_docx_path()
    doc = Document(str(template))

    hall_type = (fields.get("hall_type") or "sport").strip().lower()
    event_kind = (fields.get("event_kind") or "training").strip().lower()
    rebuild = bool(fields.get("rebuild_hall"))
    eq_volley = bool(fields.get("equipment_volleyball", True))

    today = date.today()
    mapping = {
        "{{ORG_NAME}}": (fields.get("org_name") or TROYAN_ORG_NAME).strip(),
        "{{EIK}}": (fields.get("eik") or TROYAN_EIK).strip(),
        "{{CITY}}": (fields.get("city") or TROYAN_CITY).strip(),
        "{{MUNICIPALITY}}": (fields.get("municipality") or TROYAN_MUNICIPALITY).strip(),
        "{{OBLAST}}": (fields.get("oblast") or TROYAN_OBLAST).strip(),
        "{{STREET}}": (fields.get("street") or TROYAN_STREET).strip(),
        "{{PHONE}}": (fields.get("phone") or "").strip() or "—",
        "{{REP_NAME}}": (fields.get("representative_name") or "").strip() or "—",
        "{{EMAIL}}": (fields.get("email") or "").strip() or "—",
        "{{EVENT_TITLE}}": (
            fields.get("event_title") or event_title_for_month(year, month)
        ).strip(),
        "{{TODAY}}": f"{today.day:02d}.{today.month:02d}.{today.year}",
        "{{CHK_HALL_SPORT}}": _chk(hall_type in ("sport", "sports", "спортна")),
        "{{CHK_HALL_TRAIN}}": _chk(hall_type in ("train", "training", "тренировъчна")),
        "{{CHK_HALL_CONF}}": _chk(hall_type in ("conf", "conference", "конферентна")),
        "{{CHK_KIND_TRAIN}}": _chk(event_kind in ("training", "train", "тренировки")),
        "{{CHK_KIND_TOURN}}": _chk(event_kind in ("tournament", "tourn", "турнир")),
        "{{CHK_KIND_OTHER}}": _chk(event_kind in ("other", "друго")),
        "{{CHK_EQ_BASKET}}": _chk(False),
        "{{CHK_EQ_VOLLEY}}": _chk(eq_volley),
        "{{CHK_EQ_TENNIS}}": _chk(False),
        "{{CHK_EQ_VINYL}}": _chk(False),
        "{{CHK_EQ_SOUND}}": _chk(False),
        "{{CHK_EQ_SCREEN}}": _chk(False),
        "{{CHK_REBUILD_YES}}": _chk(rebuild),
        "{{CHK_REBUILD_NO}}": _chk(not rebuild),
    }

    for paragraph in _iter_all_paragraphs(doc):
        _replace_in_paragraph(paragraph, mapping)

    # Fill schedule table (first 4-column table with Дата header)
    sched = None
    for table in doc.tables:
        if len(table.columns) == 4 and table.rows and "Дата" in (table.rows[0].cells[0].text or ""):
            sched = table
            break
    if sched is not None:
        # Ensure enough rows: header + len(slots)
        needed = max(len(slots), 10)
        while len(sched.rows) - 1 < needed:
            # clone last row
            tbl = sched._tbl
            last = sched.rows[-1]._tr
            tbl.append(deepcopy(last))

        def _set_cell(cell, text: str, *, checkbox: bool = False) -> None:
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run("☑" if checkbox else (text or ""))
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(9)

        for i in range(1, len(sched.rows)):
            slot = slots[i - 1] if i - 1 < len(slots) else None
            if slot:
                _set_cell(sched.rows[i].cells[0], slot.get("date_bg") or slot.get("date") or "")
                _set_cell(sched.rows[i].cells[1], slot.get("start_time") or "")
                _set_cell(sched.rows[i].cells[2], slot.get("end_time") or "")
                _set_cell(sched.rows[i].cells[3], "☐", checkbox=False)
            else:
                _set_cell(sched.rows[i].cells[0], "")
                _set_cell(sched.rows[i].cells[1], "")
                _set_cell(sched.rows[i].cells[2], "")
                _set_cell(sched.rows[i].cells[3], "☐")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# Backwards-compatible alias (UI/API now serve DOCX).
def build_troyan_hall_application_pdf(*, year: int, month: int, slots: list[dict[str, Any]], fields: dict[str, Any]) -> bytes:
    return build_troyan_hall_application_docx(year=year, month=month, slots=slots, fields=fields)
